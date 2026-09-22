"""Portable memo exports. User text is data in HTML, Markdown and Excel."""
from __future__ import annotations

import html
import io
import json
import re

from .analytics import snapshot, valid_evidence


def _text(value):
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return str(value)


def _markdown(value):
    value = html.escape(_text(value), quote=False)
    return re.sub(r"([\\`*_{}\[\]()#+.!|>~-])", r"\\\1", value)


def _spreadsheet(value):
    if value is None or isinstance(value, (int, float, bool)):
        return value
    value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", _text(value))
    # Leading whitespace is considered too: Excel/CSV consumers may trim it.
    if value.lstrip().startswith(("=", "+", "-", "@")) or value.startswith(("\t", "\r", "\n")):
        value = "'" + value
    return value[:32767]


def _selected_works(report, works):
    if report.get("stale"):
        raise ValueError("报告已过期；请重新生成后再导出")
    expected = {s["workId"]: s for s in report.get("comparison", {}).get("sourceSnapshot", [])}
    selected = [w for w in works if w["id"] in expected]
    actual = {s["workId"]: s for s in snapshot(selected)}
    if set(actual) != set(expected) or any(actual[wid] != expected[wid] for wid in expected):
        raise ValueError("报告来源与当前作品版本或标签不一致；请重新生成后再导出")
    return selected


def _evidence(report, works):
    ids = {eid for finding in report.get("findings", []) for eid in finding.get("evidenceIds", [])}
    return [dict(e, workTitle=w.get("title", "")) for w in works for eid, e in valid_evidence(w).items() if eid in ids]


def _md(report, works):
    comparison = report["comparison"]
    lines = ["# HerLens 研究备忘录", "", _markdown(report.get("question")), "",
             f"生成时间：{_markdown(report.get('createdAt'))} · 数据修订：{report.get('revision')} · 样本数：{report.get('sampleCount')}", "",
             "运行模式：deterministic-research（固定边界研究流程；发现依据当前 SQL 结果和原文证据生成）。", ""]
    if report.get("refusal"):
        lines += ["未生成发现：" + _markdown(report["refusal"]), ""]
    for finding in report.get("findings", []):
        lines += ["## " + _markdown(finding["title"]), "", _markdown(finding["observation"]), ""]
        for label, key in [("适用范围", "scope"), ("局限", "limitation"), ("下一步", "nextStep"), ("验收判据", "criterion"), ("证据 ID", "evidenceIds"), ("查询 ID", "queryIds")]:
            lines += [f"{label}：{_markdown(finding.get(key))}", ""]
    lines += ["## 原文证据", ""]
    for evidence in _evidence(report, works):
        lines += [f"- {_markdown(evidence['id'])} · {_markdown(evidence['workTitle'])} · {_markdown(evidence['versionId'])} · {_markdown(evidence['paragraphId'])} [{evidence['startChar']}, {evidence['endChar']})", "", "  " + _markdown(evidence["quote"]), ""]
    lines += ["## 查询与来源", "", "以下为实际执行的只读 SQLite 查询；纳入前校验记录见排除项。", ""]
    lines.extend("    " + line for line in comparison["sql"].splitlines())
    lines += ["", "参数：" + _markdown(comparison["params"]), "", "来源快照：" + _markdown(comparison["sourceSnapshot"]), "", "排除项：" + _markdown(comparison["excluded"]), "", "## 执行轨迹", ""]
    for step in report.get("trace", []):
        lines += [f"- {step['index']}. {_markdown(step['tool'])} · {_markdown(step['status'])} · {step['durationMs']} ms", "", "  输入：" + _markdown(step["input"]), "", "  输出：" + _markdown(step["output"]), ""]
    return "\n".join(lines).encode("utf-8")


def _html(report, works):
    escape = lambda v: html.escape(_text(v), quote=True)
    sections = []
    for finding in report.get("findings", []):
        details = "".join(f"<dt>{label}</dt><dd>{escape(finding.get(key))}</dd>" for label, key in
                          [("范围", "scope"), ("局限", "limitation"), ("下一步", "nextStep"), ("验收判据", "criterion"), ("证据 ID", "evidenceIds"), ("查询 ID", "queryIds")])
        sections.append(f"<section><h2>{escape(finding['title'])}</h2><p>{escape(finding['observation'])}</p><dl>{details}</dl></section>")
    evidence = "".join(f"<li><strong>{escape(e['id'])}</strong> · {escape(e['workTitle'])} · {escape(e['versionId'])} · {escape(e['paragraphId'])} [{e['startChar']}, {e['endChar']})<blockquote>{escape(e['quote'])}</blockquote></li>" for e in _evidence(report, works))
    comparison = report["comparison"]
    trace = "".join(f"<li><strong>{escape(t['tool'])}</strong> · {escape(t['status'])} · {escape(t['durationMs'])} ms<details><summary>输入与输出</summary><pre>{escape(t['input'])}</pre><pre>{escape(t['output'])}</pre></details></li>" for t in report["trace"])
    result = f"""<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<meta http-equiv="Content-Security-Policy" content="default-src 'none'; style-src 'unsafe-inline'; base-uri 'none'; form-action 'none'">
<title>HerLens 研究备忘录</title><style>
body{{max-width:900px;margin:48px auto;padding:0 24px;font:16px/1.8 system-ui,sans-serif;color:#24332d;background:#fff}}
h1,h2{{color:#245c43;line-height:1.4}}h2{{margin-top:32px}}dt{{font-weight:600}}dd{{margin:0 0 12px}}
pre{{white-space:pre-wrap;overflow-wrap:anywhere;background:#f3f6f3;padding:16px;font:12px/1.7 monospace}}
blockquote{{border-left:3px solid #6b927d;padding-left:14px;margin-left:0}}li{{overflow-wrap:anywhere}}section{{break-inside:avoid}}
@media print{{body{{margin:0;max-width:none;font-size:10pt}}details{{display:block}}details>pre{{display:block}}h2{{break-after:avoid}}}}
</style></head><body><h1>HerLens 研究备忘录</h1><p>{escape(report.get('question'))}</p>
<p>生成时间：{escape(report.get('createdAt'))} · 数据修订：{escape(report.get('revision'))} · 样本数：{escape(report.get('sampleCount'))}</p>
<p>运行模式：deterministic-research（固定边界研究流程；无生成式模型撰写结论）。</p>
<p>{escape(report.get('refusal', ''))}</p>{''.join(sections)}<h2>原文证据</h2><ul>{evidence}</ul>
<h2>实际 SQL 与参数</h2><pre>{escape(comparison['sql'])}</pre><pre>{escape(comparison['params'])}</pre>
<h2>来源快照与排除项</h2><pre>{escape(comparison['sourceSnapshot'])}</pre><pre>{escape(comparison['excluded'])}</pre>
<h2>执行轨迹</h2><ol>{trace}</ol></body></html>"""
    return result.encode("utf-8")


def _xlsx(report, works):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    book = Workbook()
    book.remove(book.active)
    sheets = {name: book.create_sheet(name) for name in ("研究备忘录", "样本与标签", "指标口径", "原文证据")}

    def append(sheet, values):
        sheet.append([_spreadsheet(value) for value in values])

    memo = sheets["研究备忘录"]
    append(memo, ["类型", "标题/字段", "内容", "范围", "局限", "下一步", "验收判据", "证据 ID", "查询 ID"])
    for key in ("question", "id", "mode", "revision", "sampleCount", "createdAt", "refusal"):
        append(memo, ["报告元数据", key, report.get(key)])
    for finding in report.get("findings", []):
        append(memo, ["发现", finding["title"], finding["observation"], finding["scope"], finding["limitation"], finding["nextStep"], finding["criterion"], finding["evidenceIds"], finding["queryIds"]])
    append(memo, ["查询", "SQL", report["comparison"]["sql"]])
    append(memo, ["查询", "params", report["comparison"]["params"]])
    append(memo, ["查询", "sourceSnapshot", report["comparison"]["sourceSnapshot"]])
    for step in report["trace"]:
        append(memo, ["执行轨迹", step["tool"], step["status"], step["input"], step["output"], step["durationMs"]])
    tags = sheets["样本与标签"]
    append(tags, ["作品 ID", "标题", "版本", "合成演示", "文本范围", "标签组", "标签值", "审核状态", "理由", "证据 ID"])
    for work in works:
        for annotation in work.get("annotations") or [{}]:
            append(tags, [work["id"], work.get("title"), work.get("versionId"), work.get("isSynthetic"), work.get("scope"), annotation.get("group"), annotation.get("value"), annotation.get("reviewStatus"), annotation.get("rationale"), annotation.get("evidenceIds")])
    metrics = sheets["指标口径"]
    append(metrics, ["记录类型", "作品 ID", "观测 ID", "版本", "合成演示", "定义", "渠道", "窗口", "计数单位", "起点", "终点", "曝光", "点击", "上报 CTR（原值）", "加权 CTR（比例）", "说明"])
    for work in works:
        observations = work.get("observations")
        if observations is None:
            observations = [work["metric"]] if work.get("metric") else []
        for observation in observations:
            append(metrics, ["原始观测", work["id"], observation.get("id"), observation.get("versionId"), work.get("isSynthetic"), observation.get("definition"), observation.get("channel"), observation.get("windowSpec"), observation.get("countingUnit"), observation.get("windowStart"), observation.get("windowEnd"), observation.get("impressions"), observation.get("clicks"), observation.get("ctrReported"), None, observation.get("observationType")])
    for group in report["comparison"]["metricGroups"]:
        append(metrics, ["SQL 汇总", group["workIds"], group["key"], None, group["isSynthetic"], group["definition"], group["channel"], group["windowSpec"], group["countingUnit"], None, None, group["impressions"], group["clicks"], None, group["ctr"], "点击总和 / 曝光总和；零分母留空"])
        metrics.cell(metrics.max_row, 15).number_format = "0.00%"
    for excluded in report["comparison"]["excluded"]:
        append(metrics, ["排除", excluded["workId"], excluded.get("observationId"), None, None, None, None, None, None, None, None, None, None, None, None, excluded["reason"]])
    evidence = sheets["原文证据"]
    append(evidence, ["证据 ID", "作品 ID", "标题", "版本", "内容哈希", "来源字段", "段落", "起点（含）", "终点（不含）", "原文", "审核状态"])
    for item in _evidence(report, works):
        append(evidence, [item.get(k) for k in ("id", "workId", "workTitle", "versionId", "contentHash", "field", "paragraphId", "startChar", "endChar", "quote", "reviewStatus")])
    for sheet in sheets.values():
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="245C43")
            cell.font = Font(color="FFFFFF", bold=True)
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = 22 if column[0].column < 3 else 34
            for cell in column:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
    result = io.BytesIO()
    book.save(result)
    return result.getvalue()


def export_report(report, works, format) -> tuple[bytes, str, str]:
    selected = _selected_works(report, works)
    safe_id = re.sub(r"[^a-zA-Z0-9_-]", "", str(report.get("id", "report")))[:80] or "report"
    exporters = {"md": (_md, "text/markdown; charset=utf-8"), "html": (_html, "text/html; charset=utf-8"),
                 "xlsx": (_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    if format not in exporters:
        raise ValueError("不支持的导出格式；请选择 md、html 或 xlsx")
    operation, media_type = exporters[format]
    return operation(report, selected), media_type, f"HerLens-{safe_id}.{format}"
