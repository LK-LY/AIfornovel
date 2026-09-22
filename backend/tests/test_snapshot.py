"""Public artifacts can only contain the fixed, explicitly synthetic cohort."""
import io
import json
from pathlib import Path
import sqlite3

import pytest
from openpyxl import load_workbook

from tools import build_snapshot as builder


@pytest.fixture
def isolated_project(tmp_path, monkeypatch):
    fixture = builder.PROJECT_ROOT / "fixtures" / "demo.json"
    (tmp_path / "fixtures").mkdir()
    (tmp_path / "fixtures" / "demo.json").write_bytes(fixture.read_bytes())
    monkeypatch.setattr(builder, "PROJECT_ROOT", tmp_path)
    return tmp_path


def test_snapshot_refuses_fixture_with_real_work(isolated_project):
    fixture = isolated_project / "fixtures" / "demo.json"
    works = json.loads(fixture.read_text(encoding="utf-8"))
    works[0]["isSynthetic"] = False
    fixture.write_text(json.dumps(works), encoding="utf-8")
    with pytest.raises(ValueError, match="every fixture work must explicitly be synthetic"):
        builder.build_snapshot()
    assert not (isolated_project / "public").exists()


def test_snapshot_ignores_existing_private_database(isolated_project, monkeypatch):
    private = isolated_project / "private-user-data"
    private.mkdir()
    marker = "PRIVATE_STORY_MUST_NEVER_ENTER_PUBLIC_ARTIFACTS_9b39"
    connection = sqlite3.connect(private / "herlens.sqlite3")
    connection.execute("CREATE TABLE private_texts (body TEXT)")
    connection.execute("INSERT INTO private_texts VALUES (?)", (marker,))
    connection.commit()
    connection.close()
    monkeypatch.setenv("HERLENS_DATA_DIR", str(private))
    monkeypatch.setenv("HERLENS_API_KEY", "PRIVATE_PROVIDER_KEY_NEVER_EXPORT")
    monkeypatch.setenv("HERLENS_BASE_URL", "https://private-provider.invalid/v1")
    snapshot = builder.build_snapshot()
    serialized = json.dumps(snapshot, ensure_ascii=False)
    assert marker not in serialized
    assert "PRIVATE_PROVIDER_KEY" not in serialized
    assert "private-provider.invalid" not in serialized
    assert len(snapshot["comparisons"]) == 8
    assert len(snapshot["report"]["trace"]) == 6
    assert all(work["isSynthetic"] is True for work in snapshot["works"])
    ids = {work["id"] for work in snapshot["works"]}
    for group, comparison in snapshot["comparisons"].items():
        assert comparison["reviewedOnly"] is True
        assert comparison["group"] == group
        assert comparison["sampleCount"] == len(ids)
        assert {source["workId"] for source in comparison["sourceSnapshot"]} == ids
        assert all(source["isSynthetic"] is True for source in comparison["sourceSnapshot"])
    connection = sqlite3.connect(private / "herlens.sqlite3")
    assert connection.execute("SELECT body FROM private_texts").fetchone()[0] == marker
    assert connection.execute("SELECT count(*) FROM sqlite_master WHERE type='table'").fetchone()[0] == 1
    connection.close()


def test_snapshot_emits_all_report_formats_from_actual_result(isolated_project, monkeypatch):
    monkeypatch.setattr(builder.sys, "argv", ["build_snapshot.py"])
    assert builder.main() == 0
    output = isolated_project / "public"
    result = json.loads((output / "demo-snapshot.json").read_text(encoding="utf-8"))
    markdown = (output / "demo-memo.md").read_text(encoding="utf-8")
    html = (output / "demo-memo.html").read_text(encoding="utf-8")
    book = load_workbook(io.BytesIO((output / "demo-memo.xlsx").read_bytes()))
    assert result["report"]["comparison"]["queryId"].replace("-", "\\-") in markdown
    assert result["report"]["comparison"]["queryId"] in html
    assert "Content-Security-Policy" in html
    assert len(book.sheetnames) == 4
    assert book["研究备忘录"]["C3"].value == result["report"]["id"]
    assert all(cell.data_type != "f" for sheet in book for row in sheet for cell in row)


def test_snapshot_cli_refuses_alternate_input_or_output(isolated_project, monkeypatch):
    monkeypatch.setattr(builder.sys, "argv", ["build_snapshot.py", "private.sqlite3"])
    with pytest.raises(ValueError, match="takes no dataset or output arguments"):
        builder.main()
    assert not (isolated_project / "public").exists()
